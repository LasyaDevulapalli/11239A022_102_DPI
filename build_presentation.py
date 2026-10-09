import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Theme Colors
NAVY = RGBColor(15, 23, 42)        # #0F172A
BLUE = RGBColor(37, 99, 235)       # #2563EB
TEAL = RGBColor(13, 148, 136)      # #0D9488
DARK_TEXT = RGBColor(30, 41, 59)   # #1E293B
MUTED_TEXT = RGBColor(100, 116, 139) # #64748B
BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
CARD_BG = RGBColor(255, 255, 255)  # White
BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0
GREEN = RGBColor(16, 185, 129)     # #10B981
ROSE = RGBColor(225, 29, 72)       # #E11D48
PURPLE = RGBColor(147, 51, 234)    # #9333EA

IMG_DIR = os.path.abspath('./presentation_screenshots')

def add_header(slide, title_text, category_text="ACCREDIVAULT PROJECT PRESENTATION"):
    # Header bar background
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = NAVY
    header_box.line.fill.background()

    # Category badge
    cat_tf = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(10), Inches(0.3)).text_frame
    cat_tf.word_wrap = True
    p0 = cat_tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = TEAL

    # Main Slide Title
    title_tf = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.6)).text_frame
    title_tf.word_wrap = True
    p1 = title_tf.paragraphs[0]
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)

def add_footer(slide, slide_num):
    # Footer line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER_COLOR
    line.line.fill.background()

    # Footer text
    ft_tf = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(10), Inches(0.35)).text_frame
    p0 = ft_tf.paragraphs[0]
    p0.text = "AccrediVault — NAAC & NBA Accreditation Evidence Repository | B.Tech CSE (AI & ML) Project"
    p0.font.size = Pt(9)
    p0.font.color.rgb = MUTED_TEXT

    # Slide number
    num_tf = slide.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.0), Inches(0.35)).text_frame
    p1 = num_tf.paragraphs[0]
    p1.text = f"Slide {slide_num}"
    p1.alignment = PP_ALIGN.RIGHT
    p1.font.size = Pt(9)
    p1.font.bold = True
    p1.font.color.rgb = BLUE

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ==========================================
# SLIDE 1 — TITLE SLIDE
# ==========================================
slide1 = prs.slides.add_slide(blank_slide_layout)
# Dark Navy Title Background
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = NAVY
bg1.line.fill.background()

# Title text frame
tf1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(2.2)).text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "AccrediVault"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)

p_sub = tf1.add_paragraph()
p_sub.text = "NAAC & NBA Accreditation Evidence Management Platform"
p_sub.font.size = Pt(22)
p_sub.font.bold = True
p_sub.font.color.rgb = TEAL

p_desc = tf1.add_paragraph()
p_desc.text = "A Paperless Accreditation Evidence Repository with Digital Signature Simulation, QR Verification and SHA-256 Tamper-Evident Audit Trail"
p_desc.font.size = Pt(14)
p_desc.font.color.rgb = RGBColor(203, 213, 225)

# Student & Institution Info Card
card1 = create_card(slide1, 1.0, 3.8, 11.333, 2.9, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
meta_tf = slide1.shapes.add_textbox(Inches(1.3), Inches(4.0), Inches(10.7), Inches(2.5)).text_frame
meta_tf.word_wrap = True

p_m0 = meta_tf.paragraphs[0]
p_m0.text = "PROJECT PRESENTATION — B.TECH CSE (ARTIFICIAL INTELLIGENCE & MACHINE LEARNING)"
p_m0.font.size = Pt(11)
p_m0.font.bold = True
p_m0.font.color.rgb = TEAL

lines = [
  ("Student Name:", "[Student Name]", "Roll / Register No:", "[Roll Number / Register No.]"),
  ("Project Guide:", "[Guide Name / Supervisor]", "Department:", "Computer Science & Engineering (AI & ML)"),
  ("Institution:", "[College / Institution Name]", "Academic Year:", "2025–26")
]

for label1, val1, label2, val2 in lines:
    p_row = meta_tf.add_paragraph()
    p_row.text = f"{label1} {val1}       |       {label2} {val2}"
    p_row.font.size = Pt(12)
    p_row.font.color.rgb = RGBColor(241, 245, 249)

# ==========================================
# SLIDE 2 — INTRODUCTION
# ==========================================
slide2 = prs.slides.add_slide(blank_slide_layout)
add_header(slide2, "Introduction to AccrediVault Platform")
add_footer(slide2, 2)

intro_items = [
    ("NAAC & NBA Accreditation Standards", "Higher Education Institutions in India are periodically evaluated by NAAC and NBA accreditation bodies, requiring thousands of verified evidence documents spanning academic, research, and governance domains."),
    ("Multi-Departmental Evidence Aggregation", "Accreditation evidence originates from diverse university departments (CSE, AIML, ECE, EEE), individual faculty members, placement cells, and IQAC quality assurance coordinators."),
    ("Challenges of Manual Evidence Management", "Traditional physical filing and unmonitored digital storage lead to misplaced documents, approval delays, lack of verifiability, and vulnerability to undetected record tampering."),
    ("The AccrediVault Solution", "AccrediVault provides a centralized, paperless accreditation repository featuring role-authenticated approval workflows, digital signature simulation, public QR verification, and SHA-256 tamper-evident audit trails."),
    ("Institutional Impact", "Streamlines audit preparation, guarantees complete document chain-of-custody accountability, and enables instant accreditation readiness tracking for NAAC 7-Criteria peer team reviews.")
]

top_pos = 1.4
for title, desc in intro_items:
    create_card(slide2, 0.8, top_pos, 11.733, 0.95, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(top_pos + 0.08), Inches(11.3), Inches(0.8)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = BLUE
    
    p2 = tb.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_TEXT
    top_pos += 1.08

# ==========================================
# SLIDE 3 — PROBLEM STATEMENT
# ==========================================
slide3 = prs.slides.add_slide(blank_slide_layout)
add_header(slide3, "Problem Statement: Challenges in Traditional Accreditation Management")
add_footer(slide3, 3)

problems = [
    ("Scattered & Unstructured Evidence", "Accreditation documents are stored across isolated personal drives, physical cabinets, and unmonitored servers without uniform indexing."),
    ("Slow Manual Approval Workflows", "Physical paper routing between Faculty, HODs, and Principals introduces severe bottlenecks and risk of document loss during review."),
    ("Lack of Status Transparency", "Faculty members and IQAC coordinators lack real-time visibility into whether evidence is uploaded, pending review, approved, or rejected."),
    ("Vulnerability of Database Audit Logs", "Standard database logs can be altered directly by administrators or bad actors without leaving a cryptographic trace of unauthorized modification."),
    ("Difficulty in External Verification", "External NAAC/NBA peer teams face difficulty verifying document authenticity, sign-off timestamps, and exact approval sequence."),
    ("Unmonitored Accreditation Readiness", "Institutions lack automated criteria-wise readiness metrics (NAAC Criteria I–VII) to identify missing evidence before formal inspection.")
]

for idx, (title, desc) in enumerate(problems):
    col = idx % 2
    row = idx // 2
    x = 0.8 + col * 5.95
    y = 1.4 + row * 1.75
    
    create_card(slide3, x, y, 5.75, 1.55, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = slide3.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(5.35), Inches(1.25)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = f"❌ {title}"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = ROSE
    
    p2 = tb.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 4 — OBJECTIVES
# ==========================================
slide4 = prs.slides.add_slide(blank_slide_layout)
add_header(slide4, "Project Objectives")
add_footer(slide4, 4)

objectives = [
    ("1. Centralized Evidence Repository", "Establish an organized digital archive indexing academic evidence by department, academic year, NAAC criterion, and document type."),
    ("2. Role-Based Access Control (RBAC)", "Implement secure authentication and role-specific workspaces for Faculty, HOD, Principal, NAAC/NBA Coordinator, and Admin."),
    ("3. Paperless Approval Workflow", "Build a multi-stage approval chain (Faculty Upload → HOD Review → Principal Approval) with mandatory rejection reasons."),
    ("4. Digital Signature Simulation", "Provide simulated institutional Principal sign-off capturing signer identity, signature ref ID, timestamp, and simulation reference ID."),
    ("5. SHA-256 Tamper-Evident Audit Trail", "Construct a cryptographic hash chain linking all workflow events to detect record modification and maintain chain-of-custody."),
    ("6. QR / Verification ID Portal", "Provide a public verification portal supporting QR codes and Verification IDs to instantly validate authentic documents."),
    ("7. Accreditation Readiness Monitoring", "Enable real-time executive monitoring of NAAC 7-Criteria compliance matrix, department readiness scores, and document stats.")
]

for idx, (title, desc) in enumerate(objectives):
    col = idx % 2
    row = idx // 2
    x = 0.8 + col * 5.95
    y = 1.4 + row * 1.35
    w = 11.733 if idx == 6 else 5.75
    
    create_card(slide4, x, y, w, 1.22, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = slide4.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.1), Inches(w - 0.4), Inches(1.0)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = f"🎯 {title}"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = BLUE
    
    p2 = tb.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 5 — ROLE-BASED ACCESS CONTROL
# ==========================================
slide5 = prs.slides.add_slide(blank_slide_layout)
add_header(slide5, "Role-Based Access Control (RBAC) Matrix")
add_footer(slide5, 5)

# Add Table for RBAC
rows, cols = 6, 4
left, top, width, height = Inches(0.8), Inches(1.4), Inches(11.733), Inches(4.5)
table_shape = slide5.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

# Set Column Widths
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(2.3)
table.columns[2].width = Inches(4.533)
table.columns[3].width = Inches(2.7)

rbac_data = [
    ["User Role", "Demo User", "Authorized Responsibilities & Workspace", "Sidebar Menu Items"],
    ["Faculty", "Prof. Rajesh Kumar\n(faculty@demo.com)", "Upload new accreditation evidence files; track personal submission approval status; view personal evidence archive.", "Executive Dashboard, Upload Evidence, My Approval Status, My Evidence, QR Verification"],
    ["HOD", "Dr. Sunita Rao\n(hod@demo.com)", "Review departmental evidence; approve or reject submissions with mandatory remarks; view department evidence index.", "Executive Dashboard, Approval Workflow, Department Evidence, Evidence Repository, QR, Audit Trail"],
    ["Principal", "Dr. A. K. Sharma\n(principal@demo.com)", "Final institutional approval; reject with reason; execute simulated digital signature sign-off with metadata.", "Executive Dashboard, Approval Workflow, Evidence Repository, Audit Trail, QR Verification"],
    ["NAAC/NBA Coordinator", "Prof. Meena Reddy\n(coordinator@demo.com)", "Executive readiness monitoring; NAAC 7-Criteria matrix tracking; audit trail review; public QR verifier (Read-Only).", "Executive Dashboard, Evidence Repository, Accreditation Status, Audit Trail, QR Verification"],
    ["Admin", "System Administrator\n(admin@demo.com)", "User management (create users); department configuration; system health monitoring and log oversight.", "Admin Dashboard, Admin Management, Evidence Repository, Audit Trail, QR Verification"]
]

for r_idx, row in enumerate(rbac_data):
    for c_idx, val in enumerate(row):
        cell = table.cell(r_idx, c_idx)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10 if r_idx > 0 else 11)
        p.font.bold = (r_idx == 0 or c_idx == 0)
        p.font.color.rgb = RGBColor(255, 255, 255) if r_idx == 0 else DARK_TEXT
        
        # Style Header Row
        if r_idx == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY

# Add Note Card
create_card(slide5, 0.8, 6.0, 11.733, 0.85, bg_color=RGBColor(241, 245, 249), border_color=BORDER_COLOR)
note_tf = slide5.shapes.add_textbox(Inches(1.0), Inches(6.08), Inches(11.3), Inches(0.7)).text_frame
note_tf.word_wrap = True
p_n = note_tf.paragraphs[0]
p_n.text = "🔒 Security Note: Each authenticated role receives a dedicated workspace with role-filtered sidebar items. Backend API endpoints enforce HTTP header role authorization (x-user-role) to prevent unauthorized action execution."
p_n.font.size = Pt(11)
p_n.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 6 — SYSTEM WORKFLOW
# ==========================================
slide6 = prs.slides.add_slide(blank_slide_layout)
add_header(slide6, "Accreditation Evidence Workflow & Chain of Custody")
add_footer(slide6, 6)

steps = [
    ("Step 1: Faculty Upload", "Faculty uploads evidence with title, department, academic year, NAAC criterion, & file attachment.", "Status: Pending HOD Review"),
    ("Step 2: HOD Review", "HOD inspects departmental evidence. Approves to forward or rejects with mandatory remarks.", "Status: Pending Principal Review"),
    ("Step 3: Principal Approval", "Principal conducts final institutional review. Approves for digital signature or rejects with reason.", "Status: Pending Digital Signature"),
    ("Step 4: Digital Signature", "Simulated Class-3 Principal Digital Signature attached with Ref ID, timestamp, & serial.", "Status: Digitally Signed / Approved"),
    ("Step 5: Central Repository", "Approved document indexed in central repository with permanent Document ID & Verification ID.", "Search & Filter Ready"),
    ("Step 6: QR Verification", "Public QR Verifier resolves Verification ID or scanned QR code to authentic metadata & audit history.", "✓ Document Verified"),
    ("Step 7: SHA-256 Audit Trail", "Cryptographic SHA-256 hash chain links previous & current hashes for complete audit integrity.", "Tamper-Evident Chain")
]

for idx, (s_title, s_desc, s_stat) in enumerate(steps):
    y = 1.35 + idx * 0.78
    create_card(slide6, 0.8, y, 11.733, 0.72, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = slide6.shapes.add_textbox(Inches(1.0), Inches(y + 0.05), Inches(11.3), Inches(0.6)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = f"{s_title}   →   {s_desc}"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = BLUE
    
    p2 = tb.add_paragraph()
    p2.text = f"System Outcome: {s_stat}"
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEAL

# ==========================================
# SLIDE 7 — MAJOR MODULES
# ==========================================
slide7 = prs.slides.add_slide(blank_slide_layout)
add_header(slide7, "Major System Modules")
add_footer(slide7, 7)

modules = [
    ("1. Authentication & RBAC", "Role-authenticated login with session persistence, role identification, and route protection."),
    ("2. Evidence Submission", "Faculty upload form supporting multi-criterion, department, academic year, and document type metadata."),
    ("3. Approval Workflow", "Multi-stage approval chain (Faculty → HOD → Principal) with stage tracking steppers and rejection remarks."),
    ("4. Digital Signature Simulation", "Simulated Class-3 Principal sign-off generating signature ref IDs, timestamps, and certificate serials."),
    ("5. Evidence Repository", "Centralized searchable document archive supporting multi-filtering (department, criterion, year, status)."),
    ("6. Accreditation Monitoring", "Executive dashboard tracking overall readiness score (%), 7-Criteria matrix, and department index."),
    ("7. SHA-256 Audit Trail", "Cryptographic hash chain linking event hashes with tamper detection and log restoration features."),
    ("8. QR Document Verification", "Public document verification portal resolving Verification IDs and QR codes to authentic details."),
    ("9. Admin Management", "System administration console for user account creation, department setup, and database stats.")
]

for idx, (title, desc) in enumerate(modules):
    col = idx % 3
    row = idx // 3
    x = 0.8 + col * 3.95
    y = 1.4 + row * 1.75
    
    create_card(slide7, x, y, 3.8, 1.6, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = slide7.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.15), Inches(3.5), Inches(1.3)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    
    p2 = tb.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 8 — SYSTEM ARCHITECTURE
# ==========================================
slide8 = prs.slides.add_slide(blank_slide_layout)
add_header(slide8, "System Architecture")
add_footer(slide8, 8)

# Top Box: Presentation Layer
create_card(slide8, 0.8, 1.4, 11.733, 1.0, bg_color=RGBColor(239, 246, 255), border_color=BLUE)
tb_top = slide8.shapes.add_textbox(Inches(1.0), Inches(1.48), Inches(11.3), Inches(0.8)).text_frame
tb_top.word_wrap = True
p = tb_top.paragraphs[0]
p.text = "PRESENTATION & WORKSPACE LAYER (React 18 + Tailwind CSS + Lucide Icons)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = BLUE
p2 = tb_top.add_paragraph()
p2.text = "Faculty Workspace  |  HOD Approval Queue  |  Principal Sign-off  |  Coordinator Monitor  |  Admin Management"
p2.font.size = Pt(10.5)
p2.font.color.rgb = DARK_TEXT

# Middle Box 1: Core Subsystems
create_card(slide8, 0.8, 2.6, 11.733, 1.6, bg_color=CARD_BG, border_color=BORDER_COLOR)
tb_mid = slide8.shapes.add_textbox(Inches(1.0), Inches(2.7), Inches(11.3), Inches(1.4)).text_frame
tb_mid.word_wrap = True
p = tb_mid.paragraphs[0]
p.text = "APPLICATION LOGIC LAYER (Node.js + Express.js REST API Services)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = NAVY
p2 = tb_mid.add_paragraph()
p2.text = "• Authentication & RBAC Middleware (x-user-role validation)\n• Document Upload Service (Multer multipart handling)\n• Multi-Stage Workflow Approval Engine (HOD & Principal approval status state machine)\n• Simulated Digital Signature Execution Service (Class-3 DSC metadata stamp generator)"
p2.font.size = Pt(10.5)
p2.font.color.rgb = DARK_TEXT

# Middle Box 2: Security & Audit
create_card(slide8, 0.8, 4.4, 11.733, 1.2, bg_color=RGBColor(240, 253, 244), border_color=GREEN)
tb_sec = slide8.shapes.add_textbox(Inches(1.0), Inches(4.48), Inches(11.3), Inches(1.0)).text_frame
tb_sec.word_wrap = True
p = tb_sec.paragraphs[0]
p.text = "SECURITY, VERIFICATION & AUDIT LAYER (SHA-256 & QR Engine)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = GREEN
p2 = tb_sec.add_paragraph()
p2.text = "• Node.js Crypto SHA-256 Tamper-Evident Hash Chain Engine (prev_hash | timestamp | action | current_hash)\n• Audit Integrity Verification Service & Tamper Detection Engine\n• Public QR & Verification ID Resolution Engine (QRCodeSVG + REST verifier)"
p2.font.size = Pt(10.5)
p2.font.color.rgb = DARK_TEXT

# Bottom Box: Database Layer
create_card(slide8, 0.8, 5.8, 11.733, 1.0, bg_color=RGBColor(248, 250, 252), border_color=BORDER_COLOR)
tb_bot = slide8.shapes.add_textbox(Inches(1.0), Inches(5.88), Inches(11.3), Inches(0.8)).text_frame
tb_bot.word_wrap = True
p = tb_bot.paragraphs[0]
p.text = "DATABASE & STORAGE LAYER (SQLite3 + Local File System)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = DARK_TEXT
p2 = tb_bot.add_paragraph()
p2.text = "Relational Tables: users, departments, documents, audit_trail  |  Local Storage: /server/uploads/ (PDF/Docs)"
p2.font.size = Pt(10.5)
p2.font.color.rgb = MUTED_TEXT

# ==========================================
# SLIDE 9 — TECHNOLOGY STACK
# ==========================================
slide9 = prs.slides.add_slide(blank_slide_layout)
add_header(slide9, "Technology Stack (Implementation Verified)")
add_footer(slide9, 9)

stack_data = [
    ("Frontend Framework", "React 18, Vite, Javascript (ES6+)", "Component-driven SPA architecture"),
    ("Styling & UI Components", "Tailwind CSS, Lucide React Icons", "Glassmorphic modern responsive UI"),
    ("Backend Runtime", "Node.js, Express.js (REST API)", "Asynchronous server logic & HTTP routing"),
    ("Database Engine", "SQLite3 (`accreditation.db`)", "Embedded relational DB with promisified queries"),
    ("File Upload Handling", "Multer Middleware", "Multipart evidence document upload storage"),
    ("Security & Hashing", "Node.js `crypto` (SHA-256)", "Cryptographic hash chain computation"),
    ("QR Code Generation", "QRCodeSVG (`qrcode.react`)", "Dynamic QR code rendering for verification"),
    ("Data Visualization", "Recharts Library", "Interactive bar charts & pie charts for metrics"),
    ("UI Feedback", "Canvas Confetti (`canvas-confetti`)", "Celebration feedback on digital signing")
]

for idx, (category, tech, detail) in enumerate(stack_data):
    col = idx % 3
    row = idx // 3
    x = 0.8 + col * 3.95
    y = 1.4 + row * 1.75
    
    create_card(slide9, x, y, 3.8, 1.6, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = slide9.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.15), Inches(3.5), Inches(1.3)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = category
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = BLUE
    
    p2 = tb.add_paragraph()
    p2.text = tech
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = NAVY
    
    p3 = tb.add_paragraph()
    p3.text = detail
    p3.font.size = Pt(10)
    p3.font.color.rgb = MUTED_TEXT

# ==========================================
# SLIDE 10 — SYSTEM OUTPUTS (SCREENSHOT COLLAGE)
# ==========================================
slide10 = prs.slides.add_slide(blank_slide_layout)
add_header(slide10, "System Outputs: Core Application User Interfaces")
add_footer(slide10, 10)

outputs = [
    ("screenshot_01_login.png", "1. Glassmorphism Login Page", "Role-authenticated sign-in interface"),
    ("screenshot_02_faculty_dashboard.png", "2. Faculty Executive Dashboard", "Metrics cards, department charts, & status pie"),
    ("screenshot_04_hod_workflow.png", "3. HOD Approval Workflow Queue", "Departmental evidence review & approval queue"),
    ("screenshot_05_principal_signature.png", "4. Principal Approval & Signature", "Principal final sign-off & Digital Signature"),
    ("screenshot_07_coordinator_status.png", "5. Coordinator Accreditation Monitor", "NAAC 7-Criteria readiness matrix & completion")
]

# Layout screenshots in 2x3 grid or balanced cards
for idx, (img_name, label, caption) in enumerate(outputs):
    col = idx % 3
    row = idx // 3
    x = 0.8 + col * 3.95
    y = 1.35 + row * 2.75
    
    # Outer Card
    create_card(slide10, x, y, 3.8, 2.6, bg_color=CARD_BG, border_color=BORDER_COLOR)
    
    # Image
    img_path = os.path.join(IMG_DIR, img_name)
    if os.path.exists(img_path):
        slide10.shapes.add_picture(img_path, Inches(x + 0.1), Inches(y + 0.1), Inches(3.6), Inches(2.0))
    
    # Caption Box below image
    tb = slide10.shapes.add_textbox(Inches(x + 0.1), Inches(y + 2.12), Inches(3.6), Inches(0.45)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(10.5)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    
    p2 = tb.add_paragraph()
    p2.text = caption
    p2.font.size = Pt(9)
    p2.font.color.rgb = MUTED_TEXT

# ==========================================
# SLIDE 11 — SECURITY, AUDIT & DOCUMENT VERIFICATION
# ==========================================
slide11 = prs.slides.add_slide(blank_slide_layout)
add_header(slide11, "Security, Audit & Document Verification System")
add_footer(slide11, 11)

sec_features = [
    ("screenshot_08_audit_verified.png", "SHA-256 Tamper-Evident Audit Trail", "Audit log entries linked via cryptographic SHA-256 hashes. verifyAuditIntegrity() verifies chain continuity across all records."),
    ("screenshot_09_audit_tampered.png", "Tamper Detection & Log Restoration", "Simulated record alteration triggers hash mismatch (brokenIndex). System detects tamper & enables 1-click log restoration."),
    ("screenshot_10_qr_verified.png", "Public QR & Verification ID Portal", "Public verifier resolves Verification IDs or scanned QR codes, displaying document status, approved metadata, & audit sequence.")
]

for idx, (img_name, label, text) in enumerate(sec_features):
    x = 0.8 + idx * 3.95
    y = 1.35
    
    create_card(slide11, x, y, 3.8, 5.4, bg_color=CARD_BG, border_color=BORDER_COLOR)
    
    img_path = os.path.join(IMG_DIR, img_name)
    if os.path.exists(img_path):
        slide11.shapes.add_picture(img_path, Inches(x + 0.15), Inches(y + 0.15), Inches(3.5), Inches(2.1))
    
    tb = slide11.shapes.add_textbox(Inches(x + 0.15), Inches(y + 2.35), Inches(3.5), Inches(2.9)).text_frame
    tb.word_wrap = True
    p1 = tb.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = BLUE
    
    p2 = tb.add_paragraph()
    p2.text = text
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = DARK_TEXT

# Note at bottom
note_card = create_card(slide11, 0.8, 6.3, 11.733, 0.55, bg_color=RGBColor(241, 245, 249), border_color=BORDER_COLOR)
tb_fn = slide11.shapes.add_textbox(Inches(1.0), Inches(6.33), Inches(11.3), Inches(0.4)).text_frame
tb_fn.word_wrap = True
p_fn = tb_fn.paragraphs[0]
p_fn.text = "💡 Digital Signature Clarification: The digital signature workflow is implemented as an institutional prototype simulation (Class-3 DSC metadata stamp) for demonstration purposes."
p_fn.font.size = Pt(10)
p_fn.font.bold = True
p_fn.font.color.rgb = NAVY

# ==========================================
# SLIDE 12 — CONCLUSION & FUTURE SCOPE
# ==========================================
slide12 = prs.slides.add_slide(blank_slide_layout)
add_header(slide12, "Conclusion & Future Enhancements")
add_footer(slide12, 12)

# Left Box: Conclusion
create_card(slide12, 0.8, 1.4, 5.75, 5.4, bg_color=CARD_BG, border_color=BORDER_COLOR)
tb_con = slide12.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.35), Inches(5.0)).text_frame
tb_con.word_wrap = True

p_c0 = tb_con.paragraphs[0]
p_c0.text = "PROJECT CONCLUSION"
p_c0.font.size = Pt(14)
p_c0.font.bold = True
p_c0.font.color.rgb = BLUE

con_points = [
    ("Centralized Evidence Archive", "Successfully replaces fragmented physical files with a unified NAAC/NBA evidence repository."),
    ("Paperless Multi-Stage Workflow", "Eliminates physical routing bottlenecks through automated Faculty → HOD → Principal review queues."),
    ("Cryptographic Audit Chain", "Ensures 100% audit accountability via SHA-256 tamper-evident hash linking."),
    ("Public Document Verification", "Provides transparent verification of approved evidence using QR codes & Verification IDs."),
    ("Executive Readiness Visibility", "Enables real-time monitoring of NAAC 7-Criteria compliance and department readiness scores.")
]

for title, desc in con_points:
    p = tb_con.add_paragraph()
    p.text = f"✓ {title}: "
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    p_d = tb_con.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = DARK_TEXT

# Right Box: Future Scope
create_card(slide12, 6.783, 1.4, 5.75, 5.4, bg_color=CARD_BG, border_color=BORDER_COLOR)
tb_fut = slide12.shapes.add_textbox(Inches(6.983), Inches(1.55), Inches(5.35), Inches(5.0)).text_frame
tb_fut.word_wrap = True

p_f0 = tb_fut.paragraphs[0]
p_f0.text = "FUTURE ENHANCEMENTS & SCOPE"
p_f0.font.size = Pt(14)
p_f0.font.bold = True
p_f0.font.color.rgb = TEAL

fut_points = [
    ("Real eSign / DSC Integration", "Integrate legally binding eSign API / Class-3 USB Token signing services (eMudhra / C-DAC)."),
    ("Cloud Deployment & Encrypted Storage", "Deploy on AWS / Azure with S3 encrypted blob storage & CDN distribution."),
    ("Institutional ERP Integration", "Connect with college ERP systems (e.g. CollPoll, Fedena) for automated faculty data sync."),
    ("Automated Queue Notifications", "Implement email and WhatsApp notifications for pending HOD & Principal review queues."),
    ("AI-Powered Evidence Analytics", "Incorporate LLM-based evidence quality check & automated criterion classification.")
]

for title, desc in fut_points:
    p = tb_fut.add_paragraph()
    p.text = f"🚀 {title}: "
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEAL
    
    p_d = tb_fut.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = DARK_TEXT

# Save Presentation
output_ppt = "AccrediVault_NAAC_NBA_Accreditation_Platform.pptx"
prs.save(output_ppt)
print(f"POWERPOINT PRESENTATION SUCCESSFULLY CREATED AT: {os.path.abspath(output_ppt)}")
